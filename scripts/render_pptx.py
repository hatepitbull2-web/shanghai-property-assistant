#!/usr/bin/env python3
"""Portable, editable PPTX export for render_reports.py report-model.json.

Requires python-pptx>=1.0. Coordinates and font sizes use the model's 96-DPI
canvas, matching the PDF renderer's 0.75-point-per-pixel scale. Text remains
editable and tables are native PowerPoint tables. Speaker notes, unknown model
properties, and raw customer/inventory metadata are deliberately not exported.

Example:
  python3 scripts/render_pptx.py --model work/report-model.json --output-dir reports
"""
from __future__ import annotations

import argparse
from io import BytesIO
import json
import math
import os
from pathlib import Path
import re
import tempfile

EMU_PER_PIXEL = 9525  # 914400 EMU/inch / 96 pixels/inch
POINTS_PER_PIXEL = 0.75
COLOR = re.compile(r"^#[0-9a-fA-F]{6}$")


def finite_number(value, label, positive=False):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"{label} 必须是有限数字。")
    if positive and value <= 0:
        raise ValueError(f"{label} 必须大于 0。")
    return float(value)


def color_value(value, label):
    if not isinstance(value, str) or not COLOR.fullmatch(value):
        raise ValueError(f"{label} 必须是 #RRGGBB 颜色。")
    return value[1:]


def validate_model(model):
    if not isinstance(model, dict):
        raise ValueError("model 必须是 JSON 对象。")
    width = finite_number(model.get("width"), "width", True)
    height = finite_number(model.get("height"), "height", True)
    # PowerPoint limits slide dimensions to 1–56 inches.
    if not 96 <= width <= 5376 or not 96 <= height <= 5376:
        raise ValueError("幻灯片宽高须在 96 到 5376 像素之间（96 DPI）。")
    font = model.get("font", "Microsoft YaHei")
    if not isinstance(font, str) or not font.strip():
        raise ValueError("font 必须是非空字体名。")
    slides = model.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("slides 必须包含至少一页。")
    for page, slide in enumerate(slides, 1):
        if not isinstance(slide, dict):
            raise ValueError(f"第 {page} 页必须是对象。")
        color_value(slide.get("background", "#FFFFFF"), f"第 {page} 页背景")
        elements = slide.get("elements")
        if not isinstance(elements, list):
            raise ValueError(f"第 {page} 页 elements 必须是数组。")
        for index, el in enumerate(elements, 1):
            label = f"第 {page} 页元素 {index}"
            if not isinstance(el, dict) or el.get("kind") not in ("text", "table"):
                raise ValueError(f"{label} 仅支持 text 或 table，未知元素不能静默丢弃。")
            x, y = (finite_number(el.get(k), label + "." + k) for k in ("x", "y"))
            w, h = (finite_number(el.get(k), label + "." + k, True) for k in ("w", "h"))
            if x < 0 or y < 0 or x + w > width + 0.1 or y + h > height + 0.1:
                raise ValueError(f"{label} 超出幻灯片边界。")
            finite_number(el.get("size", 24), label + ".size", True)
            if el["kind"] == "text":
                if not isinstance(el.get("text"), str):
                    raise ValueError(f"{label}.text 必须是文本。")
                color_value(el.get("color", "#172E31"), label + ".color")
                continue
            rows, widths = el.get("values"), el.get("widths")
            if not isinstance(widths, list) or not widths:
                raise ValueError(f"{label}.widths 必须是非空数组。")
            for column_width in widths:
                finite_number(column_width, label + ".widths", True)
            if not math.isclose(sum(widths), w, rel_tol=0, abs_tol=0.1):
                raise ValueError(f"{label} 表格列宽合计与 w 不一致。")
            if not isinstance(rows, list) or not rows:
                raise ValueError(f"{label}.values 必须是非空二维数组。")
            if any(not isinstance(row, list) or len(row) != len(widths) for row in rows):
                raise ValueError(f"{label} 表格每行必须与列数一致，避免丢失单元格。")
            for row in rows:
                for value in row:
                    if value is not None and not isinstance(value, (str, int, float, bool)):
                        raise ValueError(f"{label} 单元格只接受文本、数字、布尔值或 null。")
                    if isinstance(value, float) and not math.isfinite(value):
                        raise ValueError(f"{label} 单元格数字必须有限。")
            row_height = finite_number(el.get("row_height"), label + ".row_height", True)
            if not math.isclose(len(rows) * row_height, h, rel_tol=0, abs_tol=0.1):
                raise ValueError(f"{label} 表格行高合计与 h 不一致。")
    return model


def build_presentation(model):
    """Return a python-pptx Presentation without writing to disk."""
    validate_model(model)
    try:
        from importlib.metadata import version
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.oxml.xmlchemy import OxmlElement
        from pptx.oxml.ns import qn
        from pptx.util import Pt
    except ImportError as exc:
        raise RuntimeError("缺少 python-pptx；请运行 python3 -m pip install 'python-pptx>=1.0'。") from exc
    if int(version("python-pptx").split(".")[0]) < 1:
        raise RuntimeError("需要 python-pptx>=1.0，请先升级依赖。")

    def px(value):
        return round(float(value) * EMU_PER_PIXEL)

    def rgb(value):
        return RGBColor.from_string(color_value(value, "颜色"))

    presentation = Presentation()
    presentation.slide_width = px(model["width"])
    presentation.slide_height = px(model["height"])
    presentation.core_properties.title = "上海新房客户选房方案" + ("（演示）" if model.get("simulation") else "")
    presentation.core_properties.author = "上海房产销售助理"
    presentation.core_properties.subject = "客户房源比较与待核实事项"
    presentation.core_properties.keywords = ""
    presentation.core_properties.comments = ""
    font_name = model.get("font", "Microsoft YaHei")

    def populate(frame, content, size, color, bold=False, leading=1.30, vertical=MSO_ANCHOR.TOP):
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = vertical
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for index, line in enumerate(str(content).split("\n")):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = float(leading)
            run = paragraph.add_run()
            run.text = line
            run.font.name = font_name
            run.font.size = Pt(float(size) * POINTS_PER_PIXEL)
            run.font.bold = bool(bold)
            run.font.color.rgb = rgb(color)
            # font.name covers Latin. Set East Asian font explicitly so Chinese
            # characters do not unexpectedly inherit the Office theme font.
            properties = run._r.get_or_add_rPr()
            properties.set("lang", "zh-CN")
            for tag in ("a:ea", "a:cs"):
                node = properties.find(qn(tag))
                if node is None:
                    node = OxmlElement(tag)
                    properties.append(node)
                node.set("typeface", font_name)

    for item in model["slides"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(item.get("background", "#FFFFFF"))
        # Do not access slide.notes_slide: even a private original model must not
        # cause notes, customer objects or source metadata to enter the PPTX.
        for el in item["elements"]:
            bounds = tuple(px(el[key]) for key in ("x", "y", "w", "h"))
            if el["kind"] == "text":
                shape = slide.shapes.add_textbox(*bounds)
                populate(shape.text_frame, el["text"], el.get("size", 24), el.get("color", "#172E31"), el.get("bold", False))
                continue
            values, widths = el["values"], el["widths"]
            table = slide.shapes.add_table(len(values), len(widths), *bounds).table
            table.first_row = False
            table.first_col = False
            table.horz_banding = False
            table.vert_banding = False
            for index, width in enumerate(widths):
                table.columns[index].width = px(width)
            for ri, row in enumerate(values):
                table.rows[ri].height = px(el["row_height"])
                for ci, value in enumerate(row):
                    cell = table.cell(ri, ci)
                    fill = "#172E31" if ri == 0 else ("#E5EBE4" if ci == 0 else ("#FFFFFF" if ri % 2 else "#ECEFE9"))
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(fill)
                    populate(cell.text_frame, "" if value is None else str(value), el.get("size", 24),
                             "#FFFFFF" if ri == 0 else "#172E31", ri == 0, 1.02, MSO_ANCHOR.MIDDLE)
                    cell.margin_left = cell.margin_right = px(12)
                    cell.margin_top = px(6)
                    cell.margin_bottom = px(4)
                    properties = cell._tc.get_or_add_tcPr()
                    for side in ("lnL", "lnR", "lnT", "lnB"):
                        tag = "a:" + side
                        line = properties.find(qn(tag))
                        if line is not None:
                            properties.remove(line)
                        line = OxmlElement(tag)
                        line.set("w", "6350")  # 0.5 pt
                        solid = OxmlElement("a:solidFill")
                        shade = OxmlElement("a:srgbClr")
                        shade.set("val", "F5F3ED")
                        solid.append(shade)
                        line.append(solid)
                        properties.append(line)
    return presentation


def render_pptx(model, output):
    """Write atomically after a full model validation and return export metadata."""
    target = Path(output)
    if target.suffix.lower() != ".pptx":
        raise ValueError("输出文件须使用 .pptx 扩展名。")
    presentation = build_presentation(model)
    content = BytesIO()
    presentation.save(content)
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(prefix=".property-pptx-", suffix=".tmp", dir=target.parent, delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(content.getvalue())
        temporary.replace(target)
    finally:
        if temporary is not None and temporary.exists():
            temporary.unlink()
    return {"pptx": str(target.resolve()), "slides": len(model["slides"]),
            "native_tables": sum(e["kind"] == "table" for s in model["slides"] for e in s["elements"]),
            "speaker_notes_exported": False}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model", required=True, help="render_reports.py 输出的 report-model.json")
    destination = parser.add_mutually_exclusive_group(required=True)
    destination.add_argument("--output", help="目标 .pptx 路径")
    destination.add_argument("--output-dir", help="输出目录，依据 simulation 标记命名")
    args = parser.parse_args(argv)
    try:
        source = Path(args.model)
        model = json.loads(source.read_text(encoding="utf-8"))
        validate_model(model)
        filename = "上海新房_客户选房方案" + ("_演示" if model.get("simulation") else "") + ".pptx"
        target = Path(args.output) if args.output else Path(args.output_dir) / filename
        if target.resolve() == source.resolve() or (target.exists() and os.path.samefile(target, source)):
            raise ValueError("输出不得覆盖输入模型。")
        print(json.dumps(render_pptx(model, target), ensure_ascii=False))
    except (ValueError, OSError, RuntimeError) as exc:
        parser.error(str(exc))


if __name__ == "__main__":
    main()
