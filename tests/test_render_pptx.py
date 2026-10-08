"""Portable exporter preserves editable evidence and drops all notes/metadata."""
import copy
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("render_pptx", ROOT / "scripts/render_pptx.py")
exporter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(exporter)
HAS_PPTX = importlib.util.find_spec("pptx") is not None


def model():
    return {"width": 1280, "height": 720, "font": "Microsoft YaHei", "simulation": True,
            "raw_customer": {"内部佣金": "DO_NOT_EXPORT_PRIVATE_METADATA"},
            "slides": [{"background": "#F5F3ED", "notes": "DO_NOT_EXPORT_PRIVATE_NOTES",
                        "elements": [
                            {"kind": "text", "x": 72, "y": 48, "w": 1120, "h": 66, "text": "虚构项目核实清单", "size": 44, "color": "#172E31", "bold": True},
                            {"kind": "text", "x": 72, "y": 150, "w": 1120, "h": 90, "text": "预算口径冲突，待核实。\n通勤上限45分钟，当前未核实。", "size": 28, "color": "#64706D"},
                            {"kind": "table", "x": 72, "y": 300, "w": 1000, "h": 120, "widths": [350, 650], "row_height": 60, "size": 24,
                             "values": [["来源", "参考总价 / 万元"], ["列表与详情冲突", "1100–1400 / 700–900"]]}
                        ]}]}


class ModelValidationTests(unittest.TestCase):
    def test_unknown_element_rejected_instead_of_dropped(self):
        data = model()
        data["slides"][0]["elements"][0]["kind"] = "image"
        with self.assertRaises(ValueError):
            exporter.validate_model(data)

    def test_ragged_native_table_rejected(self):
        data = model()
        data["slides"][0]["elements"][-1]["values"][1].append("extra evidence")
        with self.assertRaises(ValueError):
            exporter.validate_model(data)

    def test_nonfinite_or_out_of_canvas_geometry_rejected(self):
        for x in (float("nan"), -2, 1290):
            data = model()
            data["slides"][0]["elements"][0]["x"] = x
            with self.assertRaises(ValueError):
                exporter.validate_model(data)


@unittest.skipUnless(HAS_PPTX, "python-pptx>=1.0 required for exporter integration tests")
class EditableExportTests(unittest.TestCase):
    def test_editable_text_native_table_and_no_notes(self):
        from pptx import Presentation
        data = model()
        before = copy.deepcopy(data)
        deck = exporter.build_presentation(data)
        buffer = BytesIO()
        deck.save(buffer)
        with ZipFile(BytesIO(buffer.getvalue())) as package:
            self.assertFalse(any(name.startswith("ppt/notesSlides/") for name in package.namelist()))
            xml = "\n".join(package.read(name).decode("utf-8", errors="replace") for name in package.namelist() if name.endswith(".xml"))
            self.assertNotIn("DO_NOT_EXPORT", xml)
            self.assertIn("a:tbl", xml)
            self.assertIn('typeface="Microsoft YaHei"', xml)
        opened = Presentation(BytesIO(buffer.getvalue()))
        self.assertEqual(opened.slide_width, 12192000)
        self.assertEqual(opened.slide_height, 6858000)
        self.assertEqual(len(opened.slides), 1)
        native = [s.table for s in opened.slides[0].shapes if s.has_table]
        self.assertEqual(len(native), 1)
        self.assertEqual(native[0].cell(1, 1).text, "1100–1400 / 700–900")
        native[0].cell(1, 1).text = "可编辑核实结果"
        self.assertEqual(native[0].cell(1, 1).text, "可编辑核实结果")
        self.assertEqual(data, before)

    def test_cli_output_directory_and_input_preserved(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            source = tmp / "model.json"
            source.write_text(json.dumps(model(), ensure_ascii=False), encoding="utf-8")
            before = source.read_bytes()
            result = subprocess.run([sys.executable, str(ROOT / "scripts/render_pptx.py"), "--model", str(source), "--output-dir", str(tmp / "reports")], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            receipt = json.loads(result.stdout)
            self.assertEqual(receipt["native_tables"], 1)
            self.assertFalse(receipt["speaker_notes_exported"])
            self.assertTrue(Path(receipt["pptx"]).is_file())
            self.assertEqual(source.read_bytes(), before)

    def test_cli_refuses_input_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "model.json"
            source.write_text(json.dumps(model()), encoding="utf-8")
            before = source.read_bytes()
            result = subprocess.run([sys.executable, str(ROOT / "scripts/render_pptx.py"), "--model", str(source), "--output", str(source)], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(source.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
